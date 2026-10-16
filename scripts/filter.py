"""Opt-in filters over data/jobs.jsonl. A filter applies ONLY when its flag is given; no flags = window + band only.

  python scripts/filter.py --visa --language --years --keyword "ESG" --keyword "net zero"

Filters (a role must PASS every filter that is on):
  --visa         drop roles where sponsorship is excluded (UK/US/CA/AU/NZ/IE regime) or JD says no sponsorship / clearance
  --language     drop roles needing a language you lack, or in a language-barrier country
  --years        drop roles asking more years than profile.yaml (stated asks only; unstated passes)
  --degree       drop roles requiring a higher degree than held
  --pedigree     yes/no: drop roles that ask for a top-tier / Ivy / IIM-ISB-class SCHOOL (grade language like "strong academic" is not school tier)
  --cert         drop roles that require a certification not held
  --early-careers  drop roles needing future graduation / current-student status
  --keywords exact|fuzzy   role must hit >= min_hits terms from keywords.yaml (its track's list); fuzzy tolerates spelling/plural variants
  --semantic     yes/no: JD meets the semantic bar vs the matching resume (embedding similarity >= semantic_min in keywords.yaml)
  --keyword K    one-off exact phrase K (case-insensitive, whole-word) must appear in title or JD; repeat = AND
  --any-keyword K  at least one of these phrases must appear (OR)
Modifiers: --strict (RISK also fails), --tier 1 2, --firm NAME, --track A|B, --country India, --undated (keep undated reqs),
           --days N --today YYYY-MM-DD
"""
import argparse
import csv
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

import yaml
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
TEXT_FIELDS = ("title", "summary", "qualifications_required", "qualifications_preferred", "jd_text", "keywords")
HARD_VISA = {"sponsorship-not-offered", "clearance-required", "right-to-work-required"}


def countries(row):
    return {c.strip().lower() for c in re.split(r"[;/]", str(row.get("country") or "")) if c.strip()}


def held_country(row, prof):
    craw = str(row.get("country") or "").lower()
    return any(o.lower() in craw for o in prof.get("work_authorised", []))


def check(name, row, prof, strict):
    """Return None if the role passes this filter, else a short reason."""
    bad = {"BLOCKED", "RISK"} if strict else {"BLOCKED"}
    ctry = countries(row)
    if name == "visa":
        if held_country(row, prof):
            return None
        if row.get("work_auth_stated") in HARD_VISA:
            return f"JD states {row['work_auth_stated']}"
        if ctry and ctry <= set(prof.get("sponsorship_excluded_countries", [])) and row.get("work_auth_stated") != "support-offered":
            return f"{row.get('country')}: sponsorship regime"
        if strict and not (ctry & set(prof.get("sponsorship_friendly_countries", []))):
            return f"{row.get('country') or 'unknown'}: sponsorship unverified"
        return None
    if name == "language":
        if row.get("gate_language") in bad:
            return "JD requires a language not held"
        if not held_country(row, prof) and ctry and ctry <= set(prof.get("language_barrier_countries", [])):
            return f"{row.get('country')}: language-barrier country"
        return None
    if name == "years":
        lvl = row.get("gate_years")
        return f"years {lvl}: asks {row.get('years_min')}+" if lvl in bad else None
    if name == "degree":
        return f"degree: {row.get('degree_required')}" if row.get("gate_degree") in bad else None
    if name == "pedigree":  # school tier only; "strong academic" is grades, "employer:" is prior-firm pedigree
        t = [x for x in str(row.get("pedigree_terms") or "").split(", ") if x.startswith("school:") and "academic" not in x and "first" not in x]
        return "school pedigree: " + "; ".join(t) if t else None
    if name == "cert":
        return "cert required" if row.get("gate_cert") in bad else None
    if name == "early_careers":
        return "needs future graduation / student status" if row.get("gate_early_careers") in bad else None
    raise ValueError(name)


def phrase_re(k):
    return re.compile(r"(?<!\w)" + re.escape(k.strip()) + r"(?!\w)", re.I)


def haystack(row):
    return " ".join(str(row.get(f) or "") for f in TEXT_FIELDS)


def kw_hits(row, kws, mode, cutoff):
    h = haystack(row).lower()
    if mode == "exact":
        return [k for k in kws if phrase_re(k).search(h)]
    from rapidfuzz import fuzz
    # terms under 6 chars (gri, crm, esg) stay whole-word: partial_ratio would match them inside "agricultural"
    return [k for k in kws if phrase_re(k).search(h) or (len(k) >= 6 and fuzz.partial_ratio(k.lower(), h) >= cutoff)]


_SEM = {}


def semantic_score(row, cfg):
    """Mean over JD sentences of best cosine to any resume line, vs the role's best track resume. Lazy: model loads only if --semantic."""
    import numpy as np
    from sentence_transformers import SentenceTransformer
    if not _SEM:
        _SEM["m"] = SentenceTransformer(cfg["semantic_model"])
        for t, f in (("A", "Purushottam_Kumar_Digital_Strategy_AI_Transformation_Resume.md"), ("B", "Purushottam_Kumar_ESG_Sustainability_Resume.md")):
            lines = [x.strip(" -*#|") for x in (ROOT / "resumes" / f).read_text(encoding="utf-8").splitlines() if len(x.strip()) > 40]
            _SEM[t] = _SEM["m"].encode(lines, normalize_embeddings=True)
    t = row.get("best_track") or "A"
    txt = " ".join(str(row.get(k) or "") for k in ("qualifications_required", "summary", "jd_text"))
    sents = [x.strip() for x in re.split(r"[.\n•;]", txt) if 40 < len(x.strip()) < 400][:40] or [str(row.get("title"))]
    S = _SEM["m"].encode(sents, normalize_embeddings=True)
    return float((S @ _SEM[t].T).max(1).mean())


def apply(rows, a, prof):
    cfg = yaml.safe_load((ROOT / "keywords.yaml").read_text(encoding="utf-8"))
    on = [n for n in ("visa", "language", "years", "degree", "pedigree", "cert", "early_careers") if getattr(a, n)]
    kw_all = [phrase_re(k) for k in a.keyword]
    kw_any = [phrase_re(k) for k in a.any_keyword]
    today = datetime.strptime(a.today, "%Y-%m-%d").date()
    cutoff = today - timedelta(days=a.days)
    kept, dropped = [], {}

    def drop(why):
        dropped[why] = dropped.get(why, 0) + 1

    for r in rows:
        pd = r.get("posted_date")
        if pd:
            if datetime.strptime(pd, "%Y-%m-%d").date() < cutoff:
                drop("window")
                continue
        elif not a.undated:
            drop("undated")
            continue
        if a.tier and str(r["tier"]) not in a.tier:
            drop("tier"); continue
        if a.firm and a.firm.lower() not in r["firm"].lower():
            drop("firm"); continue
        if a.track and r.get("best_track") != a.track:
            drop("track"); continue
        if a.country and a.country.lower() not in str(r.get("country") or "").lower():
            drop("country"); continue
        fails = [n for n in on if check(n, r, prof, a.strict)]
        if fails:
            drop(fails[0]); continue
        if a.keywords:
            kws = cfg[a.track or r.get("best_track") or "A"]
            hits = kw_hits(r, kws, a.keywords, cfg["fuzzy_cutoff"])
            if len(hits) < cfg["min_hits"]:
                drop("keywords"); continue
            r["keyword_hits"] = ", ".join(hits)
        h = haystack(r)
        if kw_all and not all(p.search(h) for p in kw_all):
            drop("keyword"); continue
        if kw_any and not any(p.search(h) for p in kw_any):
            drop("any-keyword"); continue
        if a.semantic:
            r["semantic"] = round(semantic_score(r, cfg), 3)
            if r["semantic"] < cfg["semantic_min"]:
                drop("semantic"); continue
        r["filters_passed"] = ", ".join(on + (["keyword"] if kw_all or kw_any else []) + ([f"keywords-{a.keywords}"] if a.keywords else []) + (["semantic"] if a.semantic else [])) or "none"
        r["unverified_years"] = "yes" if r.get("gate_years") == "UNSTATED" else ""
        kept.append(r)
    return kept, dropped, on


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    for n in ("visa", "language", "years", "degree", "pedigree", "cert"):
        ap.add_argument("--" + n, action="store_true")
    ap.add_argument("--early-careers", dest="early_careers", action="store_true")
    ap.add_argument("--keywords", choices=["exact", "fuzzy"])
    ap.add_argument("--semantic", action="store_true")
    ap.add_argument("--keyword", action="append", default=[])
    ap.add_argument("--any-keyword", dest="any_keyword", action="append", default=[])
    ap.add_argument("--strict", action="store_true")
    ap.add_argument("--tier", nargs="*")
    ap.add_argument("--firm")
    ap.add_argument("--track", choices=["A", "B"])
    ap.add_argument("--country")
    ap.add_argument("--undated", action="store_true")
    ap.add_argument("--days", type=int, default=7)
    ap.add_argument("--today", default=str(date.today()))
    ap.add_argument("--out", default="filtered")  # digest name prefix
    a = ap.parse_args()

    prof = yaml.safe_load((ROOT / "profile.yaml").read_text(encoding="utf-8"))
    rows = [json.loads(l) for l in (ROOT / "data" / "jobs.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
    # Funnel: one cumulative step per ON filter, in canonical order; each step becomes its own sheet in the dated tracker.
    steps = [("All", None), ("Window+Band", {})]
    for n in ("visa", "language", "years", "degree", "pedigree", "cert", "early_careers"):
        if getattr(a, n):
            steps.append(("+" + n, {n: True}))
    if a.keywords:
        steps.append(("+keywords " + a.keywords, {"keywords": a.keywords}))
    if a.semantic:
        steps.append(("+semantic", {"semantic": True}))
    off = dict(visa=False, language=False, years=False, degree=False, pedigree=False, cert=False, early_careers=False, keywords=None, semantic=False)
    cols = ["tier", "firm", "title", "band", "location", "country", "posted_date", "gate_verdict", "unverified_years", "years_min",
            "best_track", "keyword_hits", "semantic", "degree_required", "fit_A", "fit_B", "gate_reasons", "filters_passed", "url"]
    tpath = ROOT / "trackers" / f"tracker_{a.today}.xlsx"
    if not tpath.exists():
        raise SystemExit(f"{tpath} missing: run consolidate.py --today {a.today} first")
    wb = load_workbook(tpath)
    for n in [x for x in wb.sheetnames if x == "Funnel" or re.match(r"F\d+ ", x) or x == "FINAL"]:
        del wb[n]
    summary, enabled, prev, kept, dropped, on = [], {}, None, rows, {}, []
    for k, (label, add) in enumerate(steps):
        if add is not None:
            enabled.update(add)
            kept, dropped, on = apply(rows, argparse.Namespace(**{**vars(a), **off, **enabled}), prof)
        last = k == len(steps) - 1
        kept = sorted(kept, key=lambda r: (r.get("gate_verdict") != "CLEAR", str(r["tier"]), -(max(r.get("fit_A") or 0, r.get("fit_B") or 0)))) if last             else sorted(kept, key=lambda r: (str(r["tier"]), r.get("gate_verdict") != "CLEAR"))
        ws = wb.create_sheet("FINAL" if last else f"F{k} {label}"[:31])
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
        for r in kept:
            ws.append([r.get(c, "") for c in cols])
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        for ci, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(ci)].width = 55 if c in ("gate_reasons", "url", "keyword_hits") else (30 if c in ("title", "firm", "location") else 14)
        summary.append([f"F{k}" if not last else "FINAL", label, len(kept), "" if prev is None else prev - len(kept),
                        sum(r.get("gate_verdict") == "CLEAR" for r in kept), sum(r.get("gate_verdict") == "RISK" for r in kept),
                        sum(r.get("gate_verdict") == "BLOCKED" for r in kept)])
        prev = len(kept)
    fs = wb.create_sheet("Funnel", index=wb.sheetnames.index("F0 All"))
    fs.append(["sheet", "step", "roles kept", "dropped by this step", "CLEAR", "RISK", "BLOCKED"])
    for c in fs[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for row in summary:
        fs.append(row)
    for ci, w in enumerate((8, 22, 12, 20, 8, 8, 10), 1):
        fs.column_dimensions[get_column_letter(ci)].width = w
    wb.active = wb.sheetnames.index("FINAL")
    wb.save(tpath)

    desc = ", ".join(on + ([f"keywords={a.keywords}"] if a.keywords else []) + (["semantic"] if a.semantic else []) + [f'keyword="{k}"' for k in a.keyword] + [f'any-keyword="{k}"' for k in a.any_keyword]) or "none (window + band only)"
    L = [f"# Filtered roles, {a.today}", "", f"Window {a.days}d | filters ON: {desc}{' | strict' if a.strict else ''}",
         f"Kept {len(kept)} of {len(rows)} | dropped by: " + ", ".join(f"{k}={v}" for k, v in sorted(dropped.items(), key=lambda kv: -kv[1])), ""]
    for r in kept:
        L.append(f"- **{r['gate_verdict']}**{' (years unstated)' if r['unverified_years'] else ''} | T{r['tier']} {r['firm']} | {r['title']} | "
                 f"{r['location']} | {r['posted_date'] or 'undated'} | A{r.get('fit_A')}/B{r.get('fit_B')} | [apply]({r['url']})")
    (ROOT / "digests" / f"{a.out}_{a.today}.md").write_text("\n".join(L), encoding="utf-8")
    print(f"filters ON: {desc}\nkept {len(kept)} of {len(rows)}; dropped {dropped}")
    from collections import Counter
    print("degree asked by kept roles:", dict(Counter(r.get("degree_required") or "not stated" for r in kept)))
    print(f"clear {sum(r['gate_verdict']=='CLEAR' for r in kept)}, risk {sum(r['gate_verdict']=='RISK' for r in kept)}, "
          f"years-unstated {sum(bool(r['unverified_years']) for r in kept)}")


def _selftest():
    prof = {"work_authorised": ["India"], "sponsorship_excluded_countries": ["united kingdom"],
            "language_barrier_countries": ["germany"], "sponsorship_friendly_countries": ["uae"]}
    base = {"tier": "2", "firm": "X", "title": "Consultant", "country": "India", "posted_date": "2026-10-12", "gate_verdict": "CLEAR",
            "gate_years": "CLEAR", "gate_language": "CLEAR", "jd_text": "Net-zero strategy and ESG reporting"}
    uk = dict(base, country="United Kingdom")
    de = dict(base, country="Germany")
    yr = dict(base, gate_years="BLOCKED", years_min=6)
    rows = [base, uk, de, yr]
    ns = lambda **k: argparse.Namespace(**{**dict(visa=False, language=False, years=False, degree=False, pedigree=False, cert=False,
                                                   early_careers=False, keyword=[], any_keyword=[], keywords=None, semantic=False, strict=False, tier=None, firm=None,
                                                   track=None, country=None, undated=False, days=7, today="2026-10-15"), **k})
    assert len(apply(rows, ns(), prof)[0]) == 4            # no flags = nothing filtered
    assert len(apply(rows, ns(visa=True), prof)[0]) == 3   # UK out
    assert len(apply(rows, ns(language=True), prof)[0]) == 3  # Germany out
    assert len(apply(rows, ns(years=True), prof)[0]) == 3  # 6-yr ask out
    assert len(apply(rows, ns(visa=True, language=True, years=True), prof)[0]) == 1
    assert len(apply(rows, ns(keyword=["ESG"]), prof)[0]) == 4 and not apply(rows, ns(keyword=["ES"]), prof)[0]  # whole-word
    assert len(apply(rows, ns(keyword=["esg", "net-zero"]), prof)[0]) == 4 and not apply(rows, ns(keyword=["esg", "fraud"]), prof)[0]
    assert len(apply(rows, ns(any_keyword=["fraud", "esg"]), prof)[0]) == 4
    assert not apply(rows, ns(days=1), prof)[0]            # window
    print("filter selftest OK")


if __name__ == "__main__":
    import sys
    _selftest() if "--selftest" in sys.argv else main()
